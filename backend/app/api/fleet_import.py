import io
import zipfile
from fastapi import APIRouter, Depends, Request, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import select
from app.database import get_db
from app.common.response import ok
from app.models.operations import Robot
from app.services.operations_common import require, audit
from app.api.fleet import RobotInput, add_robot

router = APIRouter(prefix="/api/fleet")
HEADERS = ["code", "name", "model", "area", "ip", "hardware", "firmware", "map_id"]


@router.get("/export")
async def export(template: bool = False, actor=Depends(require("robot:view")), db=Depends(get_db)):
    from openpyxl import Workbook
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Robots"); sheet.append(HEADERS)
    if not template:
        rows = (await db.scalars(select(Robot).where(Robot.deleted == False).limit(2000))).all()
        from app.api.operations import csv_cell
        for row in rows: sheet.append([csv_cell(getattr(row,key)) for key in HEADERS])
    output = io.BytesIO(); book.save(output)
    audit(db, actor, "robot", "下载机器人档案" if not template else "下载导入模板")
    await db.commit()
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="robots.xlsx"'})


@router.post("/import-preview")
async def preview(request: Request, actor=Depends(require("robot:import")), db=Depends(get_db)):
    from openpyxl import load_workbook
    data = await request.body()
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20*1024*1024: raise ValueError("解压数据超过 20MB")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        rows = book.active.iter_rows(values_only=True)
        if list(next(rows, [])) != HEADERS: raise ValueError("表头与模板不一致")
        existing = set((await db.scalars(select(Robot.code))).all())
        valid, errors = [], []
        for number, values in enumerate(rows, 2):
            if number > 1001: raise ValueError("每次最多导入 1000 行")
            if not any(v is not None for v in values): continue
            try:
                payload = {key: str(value) if value is not None else "" for key,value in zip(HEADERS, values)}
                if any(value.startswith("=") for value in payload.values()): raise ValueError("不允许公式单元格")
                payload['map_id'] = int(payload['map_id']) if payload['map_id'] else None
                robot = RobotInput(**payload)
                if robot.code in existing: raise ValueError("编号重复")
                existing.add(robot.code)
                valid.append(robot.model_dump())
            except (ValueError, TypeError) as exc:
                errors.append({"row":number,"reason":str(exc)[:500]})
        book.close()
        return ok({"valid":valid,"errors":errors})
    except (ValueError, OSError, zipfile.BadZipFile) as exc:
        raise HTTPException(400, "导入文件无效: " + str(exc))


class ImportInput(BaseModel):
    confirmation: str
    robots: list[RobotInput] = Field(min_length=1, max_length=1000)


@router.post("/import")
async def import_robots(body: ImportInput, actor=Depends(require("robot:import")), db=Depends(get_db)):
    if body.confirmation != "确认导入": raise HTTPException(400, "请确认导入预览")
    results = []
    for robot in body.robots:
        try:
            result = await add_robot(robot, actor, db)
            # Device keys are returned only to this authenticated importer, never written to logs.
            results.append({"code":robot.code,"result":result['data']})
        except HTTPException as exc:
            await db.rollback()
            results.append({"code":robot.code,"error":exc.detail})
    return ok(results)
