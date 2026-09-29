"""Isolated integration tests: no production DB, map assets, hardware or network callbacks."""
import asyncio
import hashlib
import json
import os
import tempfile
import time
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

temporary = tempfile.TemporaryDirectory(prefix="rcs-tests-", ignore_cleanup_errors=True)
os.environ["RCS_DATABASE_URL"] = "sqlite+aiosqlite:///" + str(Path(temporary.name)/"test.db").replace("\\", "/")
os.environ["RCS_DATA_DIR"] = str(Path(temporary.name)/"data")
os.environ["RCS_WORKER_ENABLED"] = "false"
os.environ["USE_SQLITE_DEV"] = "true"
os.environ["ENVIRONMENT"] = "development"

from fastapi.testclient import TestClient
from starlette.websockets import WebSocketDisconnect
from app.main import app
from app.database import AsyncSessionLocal, engine
assert Path(engine.url.database).resolve() == (Path(temporary.name)/"test.db").resolve(), "Refusing to run integration tests against a non-test database"
from app.models.operations import Robot, TaskRun, UpgradeDetail, Callback
from app.api.integration import signature
from app.services.operations_worker import tick, callbacks_tick
from app.services.operations_common import now
from sqlalchemy import select
from app.services.map_editor import validate_map
from app.api.user import generate_token


MAP = {"width":20,"height":20,"points":[{"code":"A","x":1,"y":1,"type":"station"},
       {"code":"C","x":5,"y":5,"type":"charger"}],"paths":[{"start":"A","end":"C"}],"areas":[]}


class OperationsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.patches = [patch('app.main.ros2_service.start'), patch('app.main.ros2_service.stop')]
        for p in cls.patches: p.start()
        cls.client = TestClient(app, client=("127.0.0.1", 50000))
        cls.client.__enter__()
        login = cls.client.post('/user/login',json={"userAccount":"admin","userPassword":"admin"}).json()
        assert login['code'] == 0, login
        cls.headers = {'Authorization':login['data']}

    @classmethod
    def tearDownClass(cls):
        cls.client.__exit__(None,None,None)
        for p in cls.patches: p.stop()

    def call(self, method, path, data=None, expected=200, **kwargs):
        response=self.client.request(method,path,json=data,headers=self.headers,**kwargs)
        self.assertEqual(response.status_code,expected,response.text)
        if expected==200:
            body=response.json();self.assertEqual(body['code'],0,body);return body['data']
        return response

    def robot(self,code):
        return self.call('POST','/api/robots',{'code':code,'name':code,'model':'AMR','firmware':'1.0.0','hardware':'1.0'})

    def telemetry(self,robot,seq=1,battery=80):
        return self.client.post(f"/api/devices/{robot['id']}/telemetry",headers={'X-Device-Key':robot['device_key']},json={'seq':seq,'battery':battery,'map_version':1})

    def published(self,name):
        row=self.call('POST','/api/maps',{'name':name,'document':MAP})
        self.call('POST',f"/api/maps/{row['id']}/publish",{'revision':0,'confirmation':name})
        return row['id']

    def test_01_authentication(self):
        self.assertEqual(self.client.get('/api/robots').status_code,401)
        self.assertEqual(self.client.get('/user/current').status_code,401)
        self.assertEqual(self.client.get('/user/current',headers=self.headers).json()['data']['username'],'admin')
        self.assertEqual(self.client.get('/api/robots',headers={'Authorization':'agv_admin_1_bad'}).status_code,401)
        response=self.client.post('/user/register',json={'userAccount':'viewer','userPassword':'viewer-pass'})
        self.assertEqual(response.json()['code'],0)
        viewer_token=generate_token('viewer')
        current=self.client.get('/user/current',headers={'Authorization':viewer_token})
        self.assertEqual(current.json()['data']['roles'],['user'])
        denied=self.client.post('/api/robots',headers={'Authorization':generate_token('viewer')},json={'code':'DENIED','name':'Denied','model':'X'})
        self.assertEqual(denied.status_code,403)
        self.assertEqual(self.client.post('/user/logout',headers={'Authorization':viewer_token}).json()['code'],0)
        self.assertEqual(self.client.get('/user/current',headers={'Authorization':viewer_token}).status_code,401)
        with patch('app.api.user.verify_password',side_effect=RuntimeError('DB unavailable')):
            response=self.client.post('/user/login',json={'userAccount':'admin','userPassword':'wrong'})
            self.assertNotEqual(response.json()['code'],0)

    def test_02_fleet_battery_alarm_command(self):
        robot=self.robot('TEST-02');identity=robot['id']
        self.call('POST',f'/api/robots/{identity}/commands/goto-charge',{'confirmation':'TEST-02','key':'off'},409)
        self.assertEqual(self.telemetry(robot,battery=15).status_code,200)
        self.assertFalse(self.telemetry(robot,battery=99).json()['data']['accepted'])
        row=self.call('GET',f'/api/robots/{identity}');self.assertEqual(row['telemetry']['battery'],15)
        alarms=self.call('GET','/api/alarms',params={'robot_id':identity})['list'];self.assertEqual(len(alarms),1)
        alarm=alarms[0]
        self.call('POST',f"/api/alarms/{alarm['id']}/action",{'action':'close','note':'done'},409)
        self.call('POST',f"/api/alarms/{alarm['id']}/action",{'action':'acknowledge','note':'confirmed'})
        self.call('POST',f"/api/alarms/{alarm['id']}/action",{'action':'close','note':'charged'})
        command=self.call('POST',f'/api/robots/{identity}/commands/goto-charge',{'confirmation':'TEST-02','key':'charge'})
        duplicate=self.call('POST',f'/api/robots/{identity}/commands/goto-charge',{'confirmation':'TEST-02','key':'charge'})
        self.assertEqual(command['id'],duplicate['id']);self.assertEqual(command['status'],'queued')
        response=self.client.post(f"/api/devices/{identity}/commands/{command['id']}/receipt",headers={'X-Device-Key':robot['device_key']},json={'status':'success','result':'charger reached'})
        self.assertEqual(response.status_code,200,response.text)
        self.call('PUT',f'/api/robots/{identity}/battery-policy',{'critical':30,'low':20,'full':95},422)
        history=self.call('GET',f'/api/robots/{identity}/battery/history');self.assertEqual(len(history),1)

    def test_03_map_versions_conflict_validation(self):
        row=self.call('POST','/api/maps',{'name':'map-conflict','document':MAP});identity=row['id']
        self.call('POST',f'/api/maps/{identity}/draft/save',{'revision':0,'document':MAP})
        self.call('POST',f'/api/maps/{identity}/draft/save',{'revision':0,'document':MAP},409)
        self.call('POST',f'/api/maps/{identity}/publish',{'revision':1,'confirmation':'map-conflict'})
        versions=self.call('GET',f'/api/maps/{identity}/versions');self.assertEqual(len(versions),1)
        self.call('POST',f'/api/maps/{identity}/rollback/1',{'revision':2,'confirmation':'map-conflict'})
        self.assertEqual(len(self.call('GET',f'/api/maps/{identity}/versions')),2)
        bad=json.loads(json.dumps(MAP));bad['areas']=[{'name':'block','polygon':[[2,2],[4,2],[4,4],[2,4]]}]
        self.assertTrue(any('禁行区' in e['message'] for e in validate_map(bad)))
        bad['points'].append({'code':'alone','x':10,'y':10})
        self.assertTrue(any(e['element']=='alone' for e in validate_map(bad)))

    def test_04_task_dispatch_receipt_kpi(self):
        map_id=self.published('task-map');robot=self.robot('TASK-ROBOT');identity=robot['id']
        self.call('PUT',f'/api/robots/{identity}',{'code':robot['code'],'name':robot['name'],'model':'AMR','map_id':map_id})
        self.telemetry(robot)
        task=self.call('POST','/api/task-history',{'name':'Delivery','map_id':map_id,'robot_id':identity,'from_point':'A','to_point':'C'})
        asyncio.run(tick())
        commands=self.client.get(f'/api/devices/{identity}/commands',headers={'X-Device-Key':robot['device_key']}).json()['data']
        command=next(c for c in commands if c['command']=='execute-task')
        self.call('POST',f'/api/maps/{map_id}/publish',{'revision':1,'confirmation':'task-map'},409)
        self.call('POST',f"/api/task-history/{task['id']}/cancel",{},409)
        response=self.client.post(f"/api/devices/{identity}/commands/{command['id']}/receipt",headers={'X-Device-Key':robot['device_key']},json={'status':'success'})
        self.assertEqual(response.status_code,200,response.text)
        history=self.call('GET','/api/task-history',params={'state':'success'})
        self.assertTrue(any(t['id']==task['id'] for t in history['list']))
        kpi=self.call('GET','/api/dashboard/kpi');self.assertEqual(len(kpi['metrics']),8)

    def test_05_wms_signature_replay_idempotency(self):
        map_id=self.published('wms-map')
        application=self.call('POST','/api/integration/apps',{'code':'WMS','name':'Warehouse','ips':['127.0.0.1']})
        payload={'name':'external','externalId':'order-1','map_id':map_id,'from_point':'A','to_point':'C'}
        body=json.dumps(payload).encode();ts=str(int(time.time()))
        def send(nonce,content=body,signed_body=None):
            headers={'Content-Type':'application/json','X-App-Code':'WMS','X-Timestamp':ts,'X-Nonce':nonce,
                     'X-Signature':signature(application['secret'],'POST','/openapi/v1/tasks','WMS',ts,nonce,content if signed_body is None else signed_body)}
            return self.client.post('/openapi/v1/tasks',content=content,headers=headers)
        first=send('nonce-0001');self.assertEqual(first.status_code,200,first.text)
        self.assertEqual(send('nonce-0001').status_code,401)
        second=send('nonce-0002');self.assertEqual(second.json()['data']['id'],first.json()['data']['id'])
        self.assertEqual(send('nonce-0003',body+b' ',body).status_code,401)
        payload['name']='changed';self.assertEqual(send('nonce-0004',json.dumps(payload).encode()).status_code,409)
        self.assertNotIn('secret',self.call('GET','/api/integration/apps')['list'][0])

    def test_06_ota_chunk_review_progress(self):
        content=b'firmware-test-bytes'
        package=self.call('POST','/api/firmware/upload',{'name':'Firmware test','version':'2.0.0','model':'AMR','size':len(content),'sha256':hashlib.sha256(content).hexdigest()})
        identity=package['id']
        response=self.client.put(f'/api/firmware/{identity}/chunks/0',content=content,headers=self.headers)
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.client.put(f'/api/firmware/{identity}/chunks/0',content=content,headers=self.headers).status_code,200)
        self.call('POST',f'/api/firmware/{identity}/complete')
        self.call('POST',f'/api/firmware/{identity}/audit',{'confirmation':'Firmware test'})
        robot=self.robot('OTA-ROBOT');self.telemetry(robot)
        upgrade=self.call('POST','/api/upgrade/tasks',{'name':'upgrade-test','confirmation':'upgrade-test','package_id':identity,'robot_ids':[robot['id']],'batch_size':1,'interval_seconds':0})
        self.call('POST',f"/api/upgrade/tasks/{upgrade['id']}/start",{'confirmation':'upgrade-test'})
        asyncio.run(tick())
        progress=self.call('GET',f"/api/upgrade/tasks/{upgrade['id']}");detail=progress['details'][0]
        self.assertEqual(detail['status'],'dispatched',progress)
        self.call('POST',f"/api/robots/{robot['id']}/commands/goto-charge",{'confirmation':'OTA-ROBOT','key':'blocked'},409)
        self.assertEqual(self.client.get(f"/api/devices/{robot['id']}/firmware/{identity}",headers={'X-Device-Key':robot['device_key']}).content,content)
        response=self.client.post(f"/api/devices/{robot['id']}/upgrade-progress",headers={'X-Device-Key':robot['device_key']},json={'command_id':detail['command_id'],'stage':'success','progress':100,'installed_version':'wrong'})
        self.assertEqual(response.status_code,409)
        response=self.client.post(f"/api/devices/{robot['id']}/upgrade-progress",headers={'X-Device-Key':robot['device_key']},json={'command_id':detail['command_id'],'stage':'success','progress':100,'installed_version':'2.0.0'})
        self.assertEqual(response.status_code,200,response.text)
        asyncio.run(tick())
        self.assertEqual(self.call('GET',f"/api/upgrade/tasks/{upgrade['id']}")['status'],'completed')
        self.assertEqual(self.call('GET',f"/api/robots/{robot['id']}")['firmware'],'2.0.0')

    def test_07_websocket_snapshot_export(self):
        with self.client.websocket_connect('/ws/monitor') as ws:
            ws.send_json({'token':self.headers['Authorization']})
            frame=ws.receive_json();self.assertEqual(frame['topic'],'robot.status');self.assertIn('robots',frame['data'])
        snapshot=self.call('GET','/api/monitor/snapshot');self.assertIn('seq',snapshot)
        response=self.client.get('/api/logs/export',headers=self.headers)
        self.assertEqual(response.status_code,200,response.text)
        self.assertIn('UTC time',response.text)
        self.assertNotIn(self.headers['Authorization'],response.text)

    def test_08_image_and_xlsx(self):
        import io
        from PIL import Image
        from openpyxl import Workbook, load_workbook
        row=self.call('POST','/api/maps',{'name':'image-map','document':MAP})
        output=io.BytesIO();Image.new('L',(100,80),255).save(output,format='PNG')
        response=self.client.post(f"/api/maps/{row['id']}/image?revision=0",content=output.getvalue(),headers=self.headers)
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['data']['draft']['width'],5)
        digest=response.json()['data']['draft']['image']
        self.assertEqual(self.client.get(f"/api/maps/{row['id']}/image/{digest}",headers=self.headers).status_code,200)
        response=self.client.get('/api/fleet/export?template=true',headers=self.headers)
        book=load_workbook(io.BytesIO(response.content));self.assertEqual(book.active.cell(1,1).value,'code')
        book.active.append(['IMPORT-1','Import robot','AMR','','127.0.0.1','1.0','1.0.0',''])
        book.active.append(['IMPORT-1','Duplicate','AMR','','invalid-ip','','',''])
        output=io.BytesIO();book.save(output)
        result=self.client.post('/api/fleet/import-preview',content=output.getvalue(),headers=self.headers).json()['data']
        self.assertEqual(len(result['valid']),1);self.assertEqual(len(result['errors']),1)
        imported=self.call('POST','/api/fleet/import',{'confirmation':'确认导入','robots':result['valid']})
        self.assertIn('device_key',imported[0]['result'])

    def test_09_callback_retry_dead_letter(self):
        from unittest.mock import AsyncMock
        with patch.dict(os.environ,{'RCS_CALLBACK_HOSTS':'wms.example.test'}):
            application=self.call('POST','/api/integration/apps',{'code':'CALLBACK','name':'Callback test','ips':['127.0.0.1'],'callback_url':'https://wms.example.test/callback'})
            map_id=self.published('callback-map')
            task=self.call('POST','/api/task-history',{'name':'Callback task','map_id':map_id,'from_point':'A','to_point':'C'})
            async def prepare():
                async with AsyncSessionLocal() as db:
                    run=await db.get(TaskRun,task['id']);run.source='CALLBACK';run.external_id='cb-1';run.status='cancelled';run.finished_at=now();await db.commit()
            asyncio.run(prepare())
            client=AsyncMock();client.__aenter__.return_value=client
            client.post.return_value.status_code=500
            with patch('app.services.operations_worker.httpx.AsyncClient',return_value=client):
                for _ in range(5):
                    async def ready():
                        async with AsyncSessionLocal() as db:
                            cb=await db.scalar(select(Callback).where(Callback.task_id==task['id']))
                            if cb: cb.next_at=now();await db.commit()
                    asyncio.run(ready());asyncio.run(callbacks_tick())
            callbacks=self.call('GET','/api/integration/callbacks')['list']
            callback=next(c for c in callbacks if c['task_id']==task['id'])
            self.assertEqual(callback['status'],'dead');self.assertEqual(callback['attempts'],5)
            self.call('POST',f"/api/integration/callbacks/{callback['id']}/retry")

    def test_10_jump_locks_and_body_limit(self):
        robot=self.robot('JUMP-ROBOT');self.telemetry(robot,seq=1,battery=90);self.telemetry(robot,seq=2,battery=5)
        alarms=self.call('GET','/api/alarms',params={'robot_id':robot['id']})['list']
        self.assertTrue(any(a['code']=='BATTERY_JUMP' for a in alarms))
        self.assertFalse(any(a['code']=='BATTERY_CRITICAL' for a in alarms))
        response=self.client.post('/openapi/v1/tasks',content=b'x'*(1048576+1))
        self.assertEqual(response.status_code,413)
        map_id=self.call('POST','/api/maps',{'name':'locked-map','document':MAP})['id']
        self.call('POST',f'/api/maps/{map_id}/lock')
        with patch.dict(os.environ,{'RCS_PERMISSIONS':json.dumps({'viewer':['map:edit']})}):
            response=self.client.post(f'/api/maps/{map_id}/draft/save',headers={'Authorization':generate_token('viewer')},json={'revision':0,'document':MAP})
            self.assertEqual(response.status_code,423)

    def test_11_upgrade_reservation_and_low_battery(self):
        content=b'low-battery-firmware'
        package=self.call('POST','/api/firmware/upload',{'name':'FW3','model':'AMR','version':'3.0.0','size':len(content),'sha256':hashlib.sha256(content).hexdigest()})
        self.client.put(f"/api/firmware/{package['id']}/chunks/0",content=content,headers=self.headers)
        self.call('POST',f"/api/firmware/{package['id']}/complete")
        self.call('POST',f"/api/firmware/{package['id']}/audit",{'confirmation':'FW3'})
        robot=self.robot('LOW-OTA');self.telemetry(robot,battery=40)
        payload={'name':'low-upgrade','confirmation':'low-upgrade','package_id':package['id'],'robot_ids':[robot['id']]}
        task=self.call('POST','/api/upgrade/tasks',payload)
        self.call('POST','/api/upgrade/tasks',payload,409)
        self.call('POST',f"/api/upgrade/tasks/{task['id']}/start",{'confirmation':'low-upgrade'})
        asyncio.run(tick())
        progress=self.call('GET',f"/api/upgrade/tasks/{task['id']}")
        self.assertEqual(progress['details'][0]['status'],'skipped')
        self.assertIn('50%',progress['details'][0]['error'])

    def test_12_websocket_auth_and_user_limit(self):
        with self.assertRaises(WebSocketDisconnect):
            with self.client.websocket_connect('/ws/monitor') as ws:
                ws.send_json({'token':'invalid'});ws.receive_json()
        with ExitStack() as stack:
            for _ in range(3):
                ws=stack.enter_context(self.client.websocket_connect('/ws/monitor'))
                ws.send_json({'token':self.headers['Authorization']});ws.receive_json()
            with self.assertRaises(WebSocketDisconnect):
                with self.client.websocket_connect('/ws/monitor') as extra:
                    extra.send_json({'token':self.headers['Authorization']});extra.receive_json()


    def test_13_localization_auth_and_offline_controls(self):
        self.assertEqual(self.client.get('/api/localization/status').status_code,401)
        snapshot=self.call('GET','/api/localization/status')
        self.assertIn('localization',snapshot)
        with self.client.websocket_connect('/api/localization/ws') as ws:
            ws.send_json({'token':self.headers['Authorization']})
            self.assertIn('localization',ws.receive_json())
        with self.client.websocket_connect('/api/localization/ws') as ws:
            ws.send_json({'token':'invalid'})
            with self.assertRaises(WebSocketDisconnect):ws.receive_json()
        self.call('POST','/api/localization/stop',expected=409)
        self.call('POST','/api/localization/release',expected=409)
        self.call('POST','/api/localization/switch',{'key':'missing','x':0,'y':0,'yaw':0},404)

    def test_14_local_ros_task_adapter_receipt_flow(self):
        from unittest.mock import AsyncMock
        from app.services.local_robot import execute
        from app.services.ros2_service import ros2_service
        from app.services.vehicle_state import VehicleState
        from app.services.fleet_service import queue_command, finish_command
        from app.models.operations import Command
        identity=self.published('ROS contract map')
        robot=self.robot('ROS-14')
        async def exercise():
            async with AsyncSessionLocal() as db:
                row=await db.get(Robot,robot['id']);row.map_id=identity;row.heartbeat=now()
                task=TaskRun(name='ROS flow',map_id=identity,from_point='A',to_point='C',robot_id=row.id,operator='admin',status='dispatched')
                db.add(task);await db.flush()
                command=await queue_command(db,row,'execute-task',{'task_id':task.id,'from_point':'A','to_point':'C'},'admin')
                state=VehicleState();state.x=state.y=1
                send=AsyncMock(return_value={'status':'running'})
                with patch.object(ros2_service,'switch',{'status':'ready','map_id':identity,'version':1}),patch.object(ros2_service,'route',send),patch('app.services.local_robot.state',state):
                    status,result=await execute(db,row,command)
                self.assertEqual(status,'running');send.assert_awaited_once()
                self.assertEqual(len(send.call_args.args[0]),18)
                await finish_command(db,row,command,status,result)
                self.assertEqual(task.status,'running')
                await finish_command(db,row,command,'success','correlated ROS completion')
                self.assertEqual(task.status,'success');self.assertIsNotNone(task.finished_at)
                row.telemetry={'capabilities':['execute-task']}
                from fastapi import HTTPException
                with self.assertRaises(HTTPException):
                    await queue_command(db,row,'reboot',{},'admin')
                await db.rollback()
        asyncio.run(exercise())

    def test_15_stop_task_waits_for_device_receipt(self):
        from app.models.operations import Command
        from app.services.fleet_service import queue_command
        robot=self.robot('STOP-TASK-15');self.telemetry(robot)
        map_id=self.published('stop-task-map')
        async def prepare():
            async with AsyncSessionLocal() as db:
                row=await db.get(Robot,robot['id']);row.map_id=map_id
                await db.commit()
        asyncio.run(prepare())
        first=self.call('POST','/api/task-history',{'name':'Before dispatch','map_id':map_id,
            'robot_id':robot['id'],'from_point':'A','to_point':'C'})
        async def queue(task_id,status):
            async with AsyncSessionLocal() as db:
                row=await db.get(Robot,robot['id']);run=await db.get(TaskRun,task_id)
                command=await queue_command(db,row,'execute-task',{'task_id':task_id},'admin')
                command.status=status;run.status='running' if status=='running' else 'dispatched'
                await db.commit()
                return command.id
        queued_id=asyncio.run(queue(first['id'],'queued'))
        result=self.call('POST',f"/api/task-history/{first['id']}/stop",{'confirmation':'Before dispatch'})
        self.assertEqual(result['stop_status'],'cancelled_before_dispatch')
        self.assertEqual(self.call('GET','/api/task-history',params={'keyword':'Before dispatch'})['list'][0]['status'],'cancelled')
        second=self.call('POST','/api/task-history',{'name':'Active stop','map_id':map_id,
            'robot_id':robot['id'],'from_point':'A','to_point':'C'})
        active_id=asyncio.run(queue(second['id'],'running'))
        self.assertEqual(self.client.post(f"/api/task-history/{second['id']}/stop",
            json={'confirmation':'Active stop'}).status_code,401)
        with patch.dict(os.environ,{'RCS_PERMISSIONS':'{"admin":["task:edit"]}'}):
            self.assertEqual(self.client.post(f"/api/task-history/{second['id']}/stop",
                headers=self.headers,json={'confirmation':'Active stop'}).status_code,403)
        stop=self.call('POST',f"/api/task-history/{second['id']}/stop",{'confirmation':'Active stop'})
        self.assertEqual(stop['stop_status'],'pending')
        self.assertEqual(stop['task']['status'],'running')
        repeat=self.call('POST',f"/api/task-history/{second['id']}/stop",{'confirmation':'Active stop'})
        self.assertEqual(repeat['command']['id'],stop['command']['id'])
        device_headers={'X-Device-Key':robot['device_key']}
        outstanding=self.client.get(f"/api/devices/{robot['id']}/commands",headers=device_headers).json()['data']
        self.assertEqual(outstanding[0]['id'],stop['command']['id'])
        self.assertEqual(outstanding[0]['command'],'emergency-stop')
        receipt=self.client.post(f"/api/devices/{robot['id']}/commands/{stop['command']['id']}/receipt",
            headers=device_headers,json={'status':'success','result':'Stopped and locked'})
        self.assertEqual(receipt.status_code,200,receipt.text)
        done=self.call('GET','/api/task-history',params={'keyword':'Active stop'})['list'][0]
        self.assertEqual(done['status'],'cancelled')
        self.assertTrue(any(item['status']=='stop_requested' for item in done['timeline']))
        self.assertEqual(self.client.post(f"/api/devices/{robot['id']}/commands/{active_id}/receipt",
            headers=device_headers,json={'status':'success','result':'late success'}).status_code,409)
        async def statuses():
            async with AsyncSessionLocal() as db:
                return (await db.get(Command,queued_id)).status,(await db.get(Command,active_id)).status
        self.assertEqual(asyncio.run(statuses()),('cancelled','cancelled'))
        other=self.robot('STOP-FAIL-15');self.telemetry(other)
        async def prepare_failure():
            async with AsyncSessionLocal() as db:
                row=await db.get(Robot,other['id']);row.map_id=map_id
                run=TaskRun(name='Stop failure',map_id=map_id,robot_id=other['id'],
                    from_point='A',to_point='C',status='running',timeline=[])
                db.add(run);await db.flush()
                original=await queue_command(db,row,'execute-task',{'task_id':run.id},'admin')
                original.status='running';await db.commit()
                return run.id
        failed_task=asyncio.run(prepare_failure())
        request=self.call('POST',f'/api/task-history/{failed_task}/stop',{'confirmation':'Stop failure'})
        response=self.client.post(f"/api/devices/{other['id']}/commands/{request['command']['id']}/receipt",
            headers={'X-Device-Key':other['device_key']},json={'status':'failed','result':'No stop acknowledgment'})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.call('GET','/api/task-history',params={'keyword':'Stop failure'})['list'][0]['status'],'running')

    def test_16_legacy_route_control_is_acknowledged(self):
        from unittest.mock import AsyncMock
        from app.services.ros2_service import ros2_service
        with (patch.object(ros2_service,'start_navigation',AsyncMock(return_value={'status':'running'})) as start,
              patch.object(ros2_service,'stop_route',AsyncMock(return_value={'navigation':{'status':'cancelled'}})) as stop):
            self.assertEqual(self.call('POST','/ros2/vehicle/0')['status'],'running')
            self.assertEqual(self.call('POST','/ros2/vehicle/1')['navigation']['status'],'cancelled')
            self.assertEqual(self.call('POST','/ros2/vehicle/2')['navigation']['status'],'cancelled')
            start.assert_awaited_once();self.assertEqual(stop.await_count,2)
        self.call('POST','/ros2/vehicle/3',expected=400)

    def test_17_new_home_and_legacy_assets(self):
        self.assertEqual(self.client.get('/',follow_redirects=False).headers['location'],
                         '/ui/operations.html')
        self.assertEqual(self.client.get('/ui/index.html',follow_redirects=False).headers['location'],
                         '/ui/operations.html')
        self.assertEqual(self.client.get('/ui/operations.html').status_code,200)
        asset = self.client.get('/src/api/ApiManager.js')
        self.assertEqual(asset.status_code, 200)
        self.assertIn('axiosClient', asset.text)

    def test_18_legacy_station_editor_preserves_history(self):
        station = {'stationName':'Replacement station','stationCode':'2','positionX':'1.5',
                   'positionY':'0','mapName':'replacement-map'}
        self.assertTrue(self.call('POST','/station/add',station))
        rows = self.call('GET','/station/querynew/2',params={'mapName':'replacement-map'})[0]
        self.assertEqual(len(rows),1)
        station_id = rows[0]['id']
        self.assertEqual(rows[0]['positionX'],'1.5')
        self.call('POST','/station/add',station,expected=409)
        self.call('POST','/station/add',{'stationName':'GPS only','stationCode':'2',
            'longitude':'117','latitude':'31'},expected=422)
        self.assertTrue(self.call('POST','/station/update',{'id':station_id,'positionX':'2'}))
        self.assertEqual(self.call('GET',f'/station/query/{station_id}')['positionX'],'2')
        from app.models.route import Route
        async def linked_route():
            async with AsyncSessionLocal() as db:
                route=Route(routeName='Station dependency',stationIds=json.dumps([station_id]),isDelete=0)
                db.add(route)
                await db.commit()
                return route.id
        route_id=asyncio.run(linked_route())
        self.call('DELETE',f'/station/delete/{station_id}',expected=409)
        async def retire_route():
            async with AsyncSessionLocal() as db:
                route=await db.get(Route,route_id)
                route.isDelete=1
                await db.commit()
        asyncio.run(retire_route())
        self.assertTrue(self.call('DELETE',f'/station/delete/{station_id}'))
        self.assertEqual(self.call('GET','/station/querynew/2',params={'mapName':'replacement-map'}),[[]])

    def test_19_production_does_not_silently_switch_to_empty_sqlite(self):
        from app.main import lifespan
        from app.config import settings
        with (patch.object(settings,'ENVIRONMENT','production'),
              patch.object(settings,'USE_SQLITE_DEV',True),
              patch.object(settings,'SECRET_KEY','unique-production-test-secret-0123456789')):
            with self.assertRaisesRegex(RuntimeError,'MySQL'):
                asyncio.run(lifespan(app).__aenter__())

    def test_20_java_password_hash_remains_usable(self):
        from app.api.user import hash_password, verify_password
        from app.config import settings
        from app.models.user import User
        self.assertTrue(verify_password('test-password',hash_password('test-password')))
        self.assertFalse(verify_password('incorrect',hash_password('test-password')))
        java_hash=hashlib.md5(b'ant-robottest-password').hexdigest()
        self.assertTrue(verify_password('test-password',java_hash))
        with patch.object(settings,'ENVIRONMENT','production'):
            self.assertFalse(verify_password('test-password',java_hash))
            self.assertFalse(verify_password('test-password','test-password'))
        async def add_legacy_user():
            async with AsyncSessionLocal() as db:
                db.add(User(userAccount='java_migrated_user',
                            userPassword=java_hash,isDelete=0))
                await db.commit()
        asyncio.run(add_legacy_user())
        response=self.client.post('/user/login',json={
            'userAccount':'java_migrated_user','userPassword':'test-password'})
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()['code'],0)

    def test_21_fresh_install_refuses_preview_database(self):
        from deploy.init_fresh_db import initialize
        from deploy.bootstrap_admin import create_admin
        with self.assertRaisesRegex(RuntimeError,'MySQL'):
            asyncio.run(initialize())
        with self.assertRaisesRegex(RuntimeError,'MySQL'):
            asyncio.run(create_admin('safe-example-password'))

if __name__=='__main__':
    unittest.main(verbosity=2)
